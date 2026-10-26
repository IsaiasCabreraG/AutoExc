import pyodbc # type: ignore
import openpyxl # type: ignore
import os
import json 
import sys
import logging
import io
import smtplib
from utils import *

# Configurar logger para este módulo
logger = logging.getLogger("AutoExcel")

MAPEO_DOCS = {
    'FT': '33: Factura Electrónica'
}

MAPEO_PAGO = {
    'CC': 'CC: Cuenta Corriente',
    'CA': 'CA: Cuenta de Ahorro',
    'VV': 'VV: Vale Vista',
    'CH': 'CH: Cheque',
    'OT': 'OT: Otro'
}
MAPEO_BANCOS = {
    '001': '001 BANCO DE CHILE',
    '012': '012 BANCO ESTADO',
    '014': '014 SCOTIABANK',
    '016': '016 BANCO DE CREDITO E INVERSIONES',
    '027': '027 BANCO CORP-BANCA',
    '028': '028 BANCO BICE',
    '037': '037 BANCO SANTANDER-SANTIAGO',
    '039': '039 BANCO ITAU',
    '049': '049 BANCO SECURITY',
    '055': '055 BANCO CONSORCIO',
    '266': '266 TRANSBANK',
    '504': '504 BBVA',
    '507': '507 BANCO DEL DESARROLLO'
}


def log_debug(categoria, mensaje):
    """Helper para logs estructurados en AutoExcel"""
    logger.debug(f"[{categoria}] {mensaje}")


def open_database():
    """Busca la configuración en el JSON y abre la conexión automáticamente."""
    log_debug("DB_CONN", "→ Iniciando conexión a base de datos...")
    
    # 1. Verificar si existe el internal/config.json
    if not os.path.exists("internal/config.json"):
        log_debug("DB_CONN", "✗ No se encontró internal/config.json")
        return None

    # 2. Leer la ruta desde el JSON
    try:
        with open("internal/config.json", "r", encoding="utf-8") as f:
            config = json.load(f)
            db_path = config.get("ruta_db")
        
        log_debug("DB_CONN", f"Ruta DB desde config: {db_path}")
    except Exception as e:
        log_debug("DB_CONN", f"✗ Error al leer internal/config.json: {e}")
        return None

    if not db_path:
        log_debug("DB_CONN", "✗ La ruta de la base de datos no está configurada")
        return None

    # 3. El Driver suele ser fijo para Access (.accdb)
    driver = '{Microsoft Access Driver (*.mdb, *.accdb)}'
    
    try:
        log_debug("DB_CONN", f"→ Intentando conectar con driver: {driver}")
        conn = pyodbc.connect(f'DRIVER={driver};DBQ={db_path};')
        log_debug("DB_CONN", "✓ Conexión establecida exitosamente")
        return conn
    except Exception as e:
        log_debug("DB_CONN", f"✗✗✗ Error de conexión: {type(e).__name__}: {e}")
        return None


def datos_egreso(cursor, NEgreso: int, año: int):
    """
    Busca datos de egreso y los retorna como una lista de diccionarios con nombres claros.
    Args:
        cursor (pyodbc.Cursor): Cursor de la base de datos.
        NEgreso (int): Número de egreso a buscar.
        año (int): Año del egreso.
    Returns:
        list: Lista de diccionarios con los datos del egreso.
    """
    log_debug("DB_QUERY", "═══ CONSULTANDO DATOS DE EGRESO ═══")
    
    # 1. Preparar el número de egreso (rellenar con ceros hasta 8 caracteres)
    CpbNum = str(NEgreso).zfill(8)
    CpbAno = str(año)
    codigos_permitidos = tuple(MAPEO_DOCS.keys())
    filtro_sql = str(codigos_permitidos).replace(",)", ")")
    
    log_debug("DB_QUERY", f"Parámetros de búsqueda:")
    log_debug("DB_QUERY", f"  • Egreso original: {NEgreso} → Formateado: {CpbNum}")
    log_debug("DB_QUERY", f"  • Año: {CpbAno}")
    log_debug("DB_QUERY", f"  • Filtro tipos doc: {filtro_sql}")

    # 2. Consulta SQL con las columnas solicitadas
    query = f"""
        SELECT 
            T2.RutAux,
            T2.NomAux, 
            T2.Email,
            T1.MovNumDocRef, 
            T1.MovTipDocRef, 
            T1.MovDebe, 
            T1.FecPag,
            T3.codBanco,
            T3.CuentaCte
        FROM (cwmovim AS T1
        LEFT JOIN cwtauxi AS T2 ON T1.CodAux = T2.CodAux)
        LEFT JOIN cwtauxpagos AS T3 ON T2.CodAux = T3.CodAux
        WHERE T1.CpbNum = ? 
        AND T1.CpbAno = ?
        AND T1.CodAux <> '0000000000'
        AND T1.MovTipDocRef IN {filtro_sql}
    """
    
    try:
        log_debug("DB_QUERY", "→ Ejecutando consulta SQL...")
        cursor.execute(query, (CpbNum, CpbAno))
        filas = cursor.fetchall()
        
        log_debug("DB_QUERY", f"✓ Consulta ejecutada: {len(filas)} filas obtenidas")
        
        # 3. Transformar a estructura de diccionarios (Clave: Valor)
        registros_mapeados = []
        for i, fila in enumerate(filas):

            rut_raw = fila[0] if fila[0] is not None else ""
            rut_limpio = rut_raw.strip().replace(".", "")
            
            registro = {
                'Rut Proovedor':   rut_limpio,
                'Razon Social':    fila[1] if fila[1] is not None else "SIN NOMBRE",
                'Email':           fila[2] if fila[2] is not None else "",
                'N documento':     fila[3],
                'Tipo Documento':  fila[4],
                'Monto Documento': fila[5] if fila[5] is not None else 0,
                'Fecha Emision':   fila[6],
                'Total Pagar':     fila[5] if fila[5] is not None else 0,
                'Cod Banco':       fila[7] if fila[7] is not None else "",
                'Cuenta Cte':      fila[8] if fila[8] is not None else "",
            }
            registros_mapeados.append(registro)
            
            if i < 3:  # Mostrar solo los primeros 3 registros en debug
                log_debug("DB_QUERY", f"  Registro {i+1}: RUT={registro['Rut Proovedor']}, "
                         f"Razón Social={registro['Razon Social'][:30]}..., "
                         f"Monto=${registro['Monto Documento']:,.0f}")
        
        if len(filas) > 3:
            log_debug("DB_QUERY", f"  ... y {len(filas) - 3} registros más")
        
        log_debug("DB_QUERY", "═══ CONSULTA FINALIZADA ═══")
        return registros_mapeados
        
    except Exception as e:
        log_debug("DB_QUERY", f"✗✗✗ Error en consulta SQL: {type(e).__name__}: {e}")
        return []


def mostrar_estructura_tabla(connection, tabla):
    """Funcion encargada de mostrar la estructura de una tabla en la base de datos.
    Args:
        connection (pyodbc.Connection): Conexion a la base de datos.
        tabla (str): Nombre de la tabla a inspeccionar.
    """
    log_debug("DB_INSPECT", f"→ Inspeccionando estructura de tabla: {tabla}")
    cursor = connection.cursor()
    
    log_debug("DB_INSPECT", f"{'Columna':<25} | {'Tipo de Dato':<15} | {'Tamaño'}")
    log_debug("DB_INSPECT", "-" * 55)
    
    # El método .columns devuelve metadatos de la tabla
    for row in cursor.columns(table=tabla):
        nombre_col = row.column_name
        tipo_dato = row.type_name
        tamano = row.column_size
        
        log_debug("DB_INSPECT", f"{nombre_col:<25} | {tipo_dato:<15} | {tamano}")


def obtener_ruta_recurso(relativa):
    """ Obtiene la ruta absoluta para recursos (iconos, plantillas, etc) """
    try:
        # PyInstaller crea una carpeta temporal y guarda la ruta en _MEIPASS
        base_path = sys._MEIPASS
        log_debug("RESOURCE", f"Ejecutando desde PyInstaller: {base_path}")
    except Exception:
        base_path = os.path.abspath(".")
        log_debug("RESOURCE", f"Ejecutando desde script: {base_path}")

    ruta_completa = os.path.join(base_path, relativa)
    log_debug("RESOURCE", f"Recurso solicitado: {relativa} → {ruta_completa}")
    
    return ruta_completa


def llenar_nomina_xlsx(datos_db, numero_nomina, año, RutEmpresa=""):
    """
    Llena la plantilla de Excel usando diccionarios mapeados.
    """
    log_debug("EXCEL", "╔═══════════════════════════════════════╗")
    log_debug("EXCEL", "║     GENERANDO ARCHIVO EXCEL         ║")
    log_debug("EXCEL", "╚═══════════════════════════════════════╝")
    
    log_debug("EXCEL", f"Parámetros:")
    log_debug("EXCEL", f"  • Nómina: {numero_nomina}")
    log_debug("EXCEL", f"  • Año: {año}")
    log_debug("EXCEL", f"  • RUT Empresa: {RutEmpresa}")
    log_debug("EXCEL", f"  • Registros a procesar: {len(datos_db)}")

    if not os.path.exists("internal/config.json"):
        log_debug("EXCEL", "✗ No se encontró internal/config.json para obtener la ruta de salida")
        return None

    try:
        with open("internal/config.json", "r", encoding="utf-8") as f:
            config = json.load(f)
            carpeta_salida = config.get("ruta_output", "")
        
        log_debug("EXCEL", f"Carpeta de salida: {carpeta_salida}")
    except Exception as e:
        log_debug("EXCEL", f"✗ Error al leer internal/config.json: {e}")
        return None

    if not carpeta_salida or not os.path.exists(carpeta_salida):
        log_debug("EXCEL", f"✗ La ruta de salida no es válida: {carpeta_salida}")
        return None

    nombre_archivo = f"Nomina_{numero_nomina}-{año}.xlsx"
    ruta_final_salida = os.path.join(carpeta_salida, nombre_archivo)
    ruta_plantilla = obtener_ruta_recurso("static/Plantilla_NominaPagoV2.xlsx")
    RutEmpresa = RutEmpresa.strip(".-")

    log_debug("EXCEL", f"Archivo destino: {nombre_archivo}")
    log_debug("EXCEL", f"Ruta completa: {ruta_final_salida}")

    try:
        log_debug("EXCEL", "→ Cargando plantilla Excel...")
        
        if not os.path.exists(ruta_plantilla):
            log_debug("EXCEL", f"✗ Plantilla no encontrada: {ruta_plantilla}")
            return None
        
        wb = openpyxl.load_workbook(ruta_plantilla)
        ws = wb.active
        log_debug("EXCEL", f"✓ Plantilla cargada: {ruta_plantilla}")
        
        fila_inicio = 11
        
        # Encabezado
        log_debug("EXCEL", "→ Escribiendo encabezado...")
        ws.cell(row=5, column=3).value = int(numero_nomina + año) # C5
        ws.cell(row=6, column=3).value = RutEmpresa     # C6
        log_debug("EXCEL", f"  • C5 (ID Nómina): {numero_nomina}{año}")
        log_debug("EXCEL", f"  • C6 (RUT): {RutEmpresa}")
        
        # Datos
        log_debug("EXCEL", f"→ Escribiendo {len(datos_db)} registros...")
        totalPago = 0
        
        for i, registro in enumerate(datos_db):
            totalPago += int(registro['Total Pagar'])
            fila_actual = fila_inicio + i
            
            # Mapeo explícito a celdas
            ws.cell(row=fila_actual, column=2).value = registro['Rut Proovedor']   # B
            ws.cell(row=fila_actual, column=3).value = registro['Razon Social']    # C
            ws.cell(row=fila_actual, column=4).value = registro['N documento']     # D
            ws.cell(row=fila_actual, column=5).value = MAPEO_DOCS.get(registro['Tipo Documento'], '99: Otro Documento')  # E
            ws.cell(row=fila_actual, column=6).value = registro['Monto Documento'] # F
            ws.cell(row=fila_actual, column=7).value = registro['Fecha Emision']   # G
            ws.cell(row=fila_actual, column=12).value = registro['Total Pagar']    # L
            ws.cell(row=fila_actual, column=13).value = MAPEO_PAGO.get('CC')       # M
            ws.cell(row=fila_actual, column=14).value = MAPEO_BANCOS.get(registro['Cod Banco'], '001 BANCO DE CHILE ') # N
            
            if i < 3:
                log_debug("EXCEL", f"  Fila {fila_actual}: {registro['Razon Social'][:30]}... - ${registro['Total Pagar']:,.0f}")
        
        if len(datos_db) > 3:
            log_debug("EXCEL", f"  ... {len(datos_db) - 3} filas más escritas")
        
        log_debug("EXCEL", f"✓ Total acumulado: ${totalPago:,.0f}")
        
        # Guardar
        log_debug("EXCEL", "→ Guardando archivo...")
        wb.save(ruta_final_salida)
        log_debug("EXCEL", f"✓✓✓ Archivo guardado exitosamente: {nombre_archivo}")
        log_debug("EXCEL", "╚═══════════════════════════════════════╝")
        
        resumen = (
            f"Plantilla Iconstruye generada con éxito\n"
            f"• Nómina: {numero_nomina}-{año}\n"
            f"• Archivo: {nombre_archivo}\n"
            f"• Registros: {len(datos_db)}\n"
            f"• Total a Pagar: ${totalPago:,.0f}\n"
        ).replace(",", ".")

        return resumen

    except PermissionError:
        log_debug("EXCEL", f"✗✗✗ ERROR: Archivo bloqueado - {nombre_archivo} está abierto")
        logger.error(f"El archivo {nombre_archivo} está abierto. Ciérralo e intenta de nuevo.")
        raise Exception("El archivo de destino está abierto. Por favor, ciérrelo.")
    
    except FileNotFoundError as e:
        log_debug("EXCEL", f"✗✗✗ ERROR: Archivo no encontrado: {e}")
        raise e
    
    except Exception as e:
        log_debug("EXCEL", f"✗✗✗ ERROR INESPERADO: {type(e).__name__}: {e}")
        import traceback
        log_debug("EXCEL", f"Traceback:\n{traceback.format_exc()}")
        raise e

def llenarR2(datos_db, numero_nomina, año, fecha_pago, user, app_pasword, correos_enviar, enviarBool):
    """
    Genera la nómina R2 en memoria. 
    Si enviarBool es True, requiere éxito en el envío de correo para guardar en disco.
    """
    log_debug("EXCEL", "╔═══════════════════════════════════════╗")
    log_debug("EXCEL", "║      GENERANDO PLANTILLA BANCO R2     ║")
    log_debug("EXCEL", "╚═══════════════════════════════════════╝")

    # 1. VALIDACIÓN INICIAL DE CREDENCIALES (Solo si se va a enviar correo)
    if enviarBool and (not user or not app_pasword):
        log_debug("EXCEL", "✗ Abortado: Se requiere envío de correo pero faltan credenciales del remitente.")
        return "[ERROR] R2 no generada: Configure el correo y clave en la pestaña de Correos."

    if not os.path.exists("internal/config.json"):
        return "[ERROR] R2 no generada: No se encontró internal/config.json"

    try:
        # 2. CARGAR CONFIGURACIÓN PARA RUTA DE SALIDA
        with open("internal/config.json", "r", encoding="utf-8") as f:
            config = json.load(f)
            carpeta_salida = config.get("ruta_output", "")
        
        if not carpeta_salida or not os.path.exists(carpeta_salida):
            return "[ERROR] R2 no generada: Ruta de salida no válida en configuración."

        # 3. AGRUPAMIENTO DE DATOS Y VALIDACIÓN DE CORREOS PROVEEDORES
        agrupados = {}
        proveedores_sin_correo = []
        emails_cargados = 0
        monto_total_nomina = 0

        for registro in datos_db:
            rut = str(registro['Rut Proovedor']).replace(".", "").replace("-", "")
            razon_social = registro['Razon Social']
            
            if rut not in agrupados:
                email_actual = registro.get('Email', '')
                if not email_actual or str(email_actual).strip() in ["S/E", "sin@email.cl", "", "None"]:
                    if razon_social not in proveedores_sin_correo:
                        proveedores_sin_correo.append(razon_social)
                else:
                    emails_cargados += 1

                agrupados[rut] = {
                    "rut": rut, "razon_social": razon_social, "cta_abono": registro['Cuenta Cte'],
                    "cod_banco": registro['Cod Banco'], "modalidad": "003", "facturas": [],
                    "monto_total_abonado": 0, "mail": email_actual 
                }
            
            num_doc_raw = str(registro['N documento'])
            num_doc_limpio = str(int(float(num_doc_raw))) if num_doc_raw.replace('.','',1).isdigit() else num_doc_raw

            agrupados[rut]["facturas"].append({"numero": num_doc_limpio, "monto": registro['Monto Documento']})
            agrupados[rut]["monto_total_abonado"] += registro['Monto Documento']
            monto_total_nomina += registro['Monto Documento']

        if proveedores_sin_correo:
            log_debug("EXCEL", f"⚠️ R2 cancelada: Faltan {len(proveedores_sin_correo)} correos")
            return f"[ERROR] R2 no generada, faltan correos para: {', '.join(proveedores_sin_correo)}"

        # 4. PROCESAR EXCEL EN MEMORIA (RAM)
        ruta_plantilla = obtener_ruta_recurso("static/R2.xlsx")
        wb = openpyxl.load_workbook(ruta_plantilla)
        ws = wb.active
        
        for i, (rut, info) in enumerate(agrupados.items()):
            fila_actual = 2 + i
            ws.cell(row=fila_actual, column=1).value = info["rut"]
            ws.cell(row=fila_actual, column=2).value = info["razon_social"]
            ws.cell(row=fila_actual, column=3).value = info["cta_abono"]
            ws.cell(row=fila_actual, column=4).value = info["cod_banco"]
            ws.cell(row=fila_actual, column=5).value = info["modalidad"]

            numeros_facturas = [f["numero"] for f in info["facturas"]]
            detalle_glosa = f"Detalle Factura : {', '.join(numeros_facturas)}."

            for idx, fac in enumerate(info["facturas"]):
                if idx < 11:
                    ws.cell(row=fila_actual, column=7+(idx*2)).value = int(float(fac["numero"]))
                    ws.cell(row=fila_actual, column=8+(idx*2)).value = fac["monto"]

            ws.cell(row=fila_actual, column=29).value = info["monto_total_abonado"]
            ws.cell(row=fila_actual, column=30).value = info["mail"]
            ws.cell(row=fila_actual, column=31).value = detalle_glosa

        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        log_debug("EXCEL", "✓ Estructura Excel generada correctamente en memoria RAM")

        # 5. LÓGICA DE ENVÍO Y GUARDADO FÍSICO
        nombre_archivo = f"Nomina_Banco_{numero_nomina}-{año}.xlsx"
        resumen_email = ""

        if enviarBool:
            log_debug("EMAIL", f"→ Iniciando protocolo de notificación (Destinatarios: {len(correos_enviar)})")
            datos_asunto = {'nomina': numero_nomina, 'año': año}
            
            exito_mail, msg_mail = enviarcorreo(user, app_pasword, correos_enviar, datos_asunto, agrupados, monto_total_nomina, fecha_pago)

            if exito_mail:
                log_debug("EMAIL", "✓ Notificación enviada con éxito. Procediendo a guardar archivo.")
                resumen_email = f"• Notificación: Enviada con éxito a {len(correos_enviar)} destinatarios\n"
            else:
                log_debug("EMAIL", f"✗ Fallo en el envío: {msg_mail}. Se aborta el guardado físico.")
                return f"[ERROR] R2 no generada: El correo de notificación falló ({msg_mail})."
        else:
            log_debug("EXCEL", "ℹ El envío de correo está desactivado. Se guardará directamente.")
            resumen_email = "• Notificación: Desactivada por el usuario\n"

        # 6. GUARDADO EN DISCO (Si llegó aquí, es porque enviarBool era False o el correo fue exitoso)
        ruta_final_salida = os.path.join(carpeta_salida, nombre_archivo)
        with open(ruta_final_salida, "wb") as f:
            f.write(excel_buffer.getbuffer())
        
        log_debug("EXCEL", f"✓ Archivo consolidado en disco: {ruta_final_salida}")

        return (
            f"Plantilla Banco generada correctamente\n"
            f"• Archivo: {nombre_archivo}\n"
            f"• Proveedores: {len(agrupados)}\n"
            f"• Emails en glosa: {emails_cargados} / {len(agrupados)}\n"
            + resumen_email
        )

    except Exception as e:
        log_debug("EXCEL", f"✗ Error crítico: {str(e)}")
        return f"[ERROR] Plantilla Banco: Fallo técnico durante el procesamiento."
    
def obtener_años_disponibles():
    """
    Intenta conectar a la DB usando la ruta de internal/config.json 
    y retorna la lista de años.
    Returns:
        list: Lista de años disponibles en la tabla cwmovim.
    """
    log_debug("DB_AÑOS", "→ Obteniendo años disponibles...")
    años = []
    conn = open_database()
    
    if conn:
        try:
            cursor = conn.cursor()
            log_debug("DB_AÑOS", "→ Ejecutando consulta DISTINCT de años...")
            cursor.execute("SELECT DISTINCT CpbAno FROM cwmovim ORDER BY CpbAno DESC")
            años = [str(fila[0]) for fila in cursor.fetchall()]
            log_debug("DB_AÑOS", f"✓ {len(años)} años encontrados: {', '.join(años)}")
        except Exception as e:
            log_debug("DB_AÑOS", f"✗ Error al obtener años: {type(e).__name__}: {e}")
        finally:
            conn.close()
            log_debug("DB_AÑOS", "Conexión cerrada")
    else:
        log_debug("DB_AÑOS", "✗ No se pudo conectar a la base de datos")
    
    return años


def inspeccionar_tipos_documento():
    """Consulta los valores únicos y la frecuencia de la columna MovTipDocRef."""
    log_debug("DB_INSPECT", "→ Inspeccionando tipos de documento...")
    conn = open_database()
    resultados = []
    
    if conn:
        try:
            cursor = conn.cursor()
            query = """
                SELECT MovTipDocRef, COUNT(*) as Total
                FROM cwmovim
                GROUP BY MovTipDocRef
                ORDER BY COUNT(*) DESC
            """
            cursor.execute(query)
            resultados = cursor.fetchall()
            
            log_debug("DB_INSPECT", f"✓ {len(resultados)} tipos de documento encontrados:")
            for t in resultados:
                log_debug("DB_INSPECT", f"  • Código: '{t[0]}' | Registros: {t[1]}")
                
        except Exception as e:
            log_debug("DB_INSPECT", f"✗ Error al inspeccionar documentos: {e}")
        finally:
            conn.close()
    
    return resultados


def enviarcorreo( remitente, password, destinatarios, asunto_data, datos_reporte, monto_total, fecha_pago):
    """
    Construye el HTML y envía el correo.
    asunto_data: dict con {'nomina': num, 'año': año}
    datos_reporte: el diccionario 'agrupados' de llenarR2
    """
    import smtplib
    from email.mime.text import MIMEText
    from email.mime.multipart import MIMEMultipart

    # 1. Construir Asunto y Cuerpo HTML
    asunto = f"Detalle Egreso N° {asunto_data['nomina']} ({asunto_data['año']}) pagado"
    
    filas_tabla = ""
    for _, info in datos_reporte.items():
        monto_fmt = f"${info['monto_total_abonado']:,.0f}".replace(",", ".")
        numeros_facturas = [f["numero"] for f in info["facturas"]]
        glosa = f"{', '.join(numeros_facturas)}."
        rut = info['rut'][:-1] + "-" + info['rut'][-1]
        filas_tabla += f"""
            <tr>
                <td style="border: 1px solid #ddd; width: 36%; padding: 4px; font-size: 13px;">{info['razon_social']}</td>
                <td style="border: 1px solid #ddd; width: 7%;padding: 4px; font-size: 13px;">{rut}</td>
                <td style="border: 1px solid #ddd; width: 7%;padding: 4px; font-size: 13px; text-align: right; font-weight: bold;">{monto_fmt}</td>
                <td style="border: 1px solid #ddd; width: 50%;padding: 4px; font-size: 12px; color: #555;">{glosa}</td>
            </tr>
        """

    cuerpo_html = f"""
    <html>
    <body style="font-family: 'Segoe UI', sans-serif; color: #333;">
        <p>De nuestra consideración:\n
    Se ha procedido a dar orden de pago a los siguientes documentos:<p> 
        <h2 style="color: #2c3e50; border-top: 2px solid #3498db; padding: 4px;">
            Resumen de Nómina - N° {asunto_data['nomina']} - fecha de abono: {fecha_pago}
        </h2>
        <div style="margin: 0 3%;" >
            <table style="border-collapse: collapse; width: 100%; margin: 10px;">
                <thead>
                    <tr style="background-color: #3498db; color: white;">
                        <th style="padding: 8px; text-align: left;">RazónsSocial</th>
                        <th style="padding: 8px; text-align: left;">RUT</th>
                        <th style="padding: 8px; text-align: right;">Monto</th>
                        <th style="padding: 8px; text-align: left;">Detalle de facturas</th>
                    </tr>
                </thead>
                <tbody>{filas_tabla}</tbody>
                <tfoot>
                    <tr style="background-color: #f9f9f9; font-weight: bold;">
                        <td colspan="2" style="padding: 8px; border-top: 2px solid #333; text-align: right;">TOTAL:</td>
                        <td style="padding: 8px; border-top: 2px solid #333; text-align: right;">${monto_total:,.0f}</td>
                        <td style="border-top: 2px solid #333;"></td>
                    </tr>
                </tfoot>
            </table>
        </div>
    </body>
    </html>
    """.replace(",", ".")

    # 2. Lógica de Envío (igual que antes pero usando destinatarios limpios)
    try:
        destinatarios_limpios = [d.strip() for d in destinatarios if d and str(d).strip()]
        if not destinatarios_limpios: return True, "Sin destinatarios"
        log_debug("EMAIL", f" destinatarios: {destinatarios_limpios}")

        mensaje = MIMEMultipart()
        mensaje["From"] = remitente
        mensaje["To"] = ", ".join(destinatarios_limpios)
        mensaje["Subject"] = asunto
        mensaje.attach(MIMEText(cuerpo_html, "html", "utf-8"))

        server = smtplib.SMTP("smtp.gmail.com", 587)
        server.starttls()
        server.login(remitente, password)
        server.send_message(mensaje)
        server.quit()
        return True, f"se enviaron {len(destinatarios_limpios)} correos."
    except Exception as e:
        return False, str(e)

        


if __name__ == "__main__":
        # Prueba rápida de conexión y obtención de años
        conn = open_database()
        if conn:
            mostrar_estructura_tabla(conn, "cwmovim")
            años = obtener_años_disponibles()
            print(f"Años disponibles en la DB: {años}")
            conn.close()
        else:
            print("No se pudo conectar a la base de datos.")    

